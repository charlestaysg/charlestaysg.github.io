"""Check the built page against the Pulse, the Landscape and the SSIC table. Exit 1 on any failure.

Run:  python3 check.py   (after ssic.py and build.py)
Sector figures are recomputed from the Pulse snapshot, course figures read from the Landscape's
latest.json, and top-course attendance counted straight from the Landscape's raw XLSX with plain
openpyxl (not through its clean.py).
"""
import ast, json, os, pathlib, re, sys

sys.dont_write_bytecode = True
HERE = pathlib.Path(__file__).parent
# refresh.sh points these at fresh copies from the site repo; by default, the vault sibling folders
def sibling(*names):  # vault folder names first, then the site repo's tools/ names
    return next((HERE.parent / n for n in names if (HERE.parent / n).exists()), HERE.parent / names[0])
PULSE = pathlib.Path(os.environ.get("UEN_PULSE_DIR") or sibling("sme-workforce-pulse", "sme-pulse"))
LAND = pathlib.Path(os.environ.get("UEN_LAND_DIR") or sibling("skills-training-landscape", "skills-landscape"))
fails = []


def check(name, ok, detail=""):
    print(f"{'PASS' if ok else 'FAIL'}  {name}{'  ' + detail if detail else ''}")
    if not ok:
        fails.append(name)


html = (HERE / "index.html").read_text()
baked = json.loads((HERE / "data" / "baked.json").read_text())
m = re.search(r"const D = (\{.*?\});\nconst fmt", html, re.S)
page = json.loads(m.group(1)) if m else {}
check("page carries baked.json unchanged", page == baked)
check("no unreplaced placeholders", "/*__" not in html)
size = len(html.encode())
check("page under 200 KB", size < 200_000, f"{size // 1024} KB")
check("no external scripts", not re.search(r"<script[^>]+src=", html))

# sectors vs the Pulse
tree = ast.parse((PULSE / "build.py").read_text())
ps = next(ast.literal_eval(n.value) for n in ast.walk(tree)
          if isinstance(n, ast.Assign) and getattr(n.targets[0], "id", "") == "SECTORS")
check("sector keys and labels match the Pulse build.py",
      [(k, v["label"]) for k, v in baked["sectors"].items()] == [(k, l) for k, l, _, _ in ps])
check("MOM names match the Pulse", all(baked["sectors"][k]["mom"] == mom for k, _, mom, _ in ps))

ds = json.loads((PULSE / "data" / "snapshot.json").read_text())["datasets"]
quarters = sorted({r["q"] for r in ds["vacancies_total"]["data"]})
q, qya = quarters[-1], quarters[-5]
check("Pulse quarter is the latest", baked["asof"]["pulse_q"] == q, q)


def val(key, industry, quarter, occ=None):
    for r in ds[key]["data"]:
        if r["q"] == quarter and r.get("industry") == industry and r.get("occupation") == occ:
            return r["v"]


bad = []
PM = "professionals, managers, executives and technicians"
for k, _, mom, _ in ps:
    p = baked["sectors"][k]["pulse"]
    exp = {"vac": val("vacancies", mom, q), "vac_ya": val("vacancies", mom, qya), "rate": val("vacancy_rate", mom, q),
           "rate_ya": val("vacancy_rate", mom, qya), "ret": val("retrenched", mom, q), "ret_ya": val("retrenched", mom, qya),
           "pmet_vac": val("vacancies_occ", mom, q, PM)}
    exp = {a: b for a, b in exp.items()}
    got = {a: p[a] for a in exp}
    r4 = [val("retrenched_occ", mom, x, PM) for x in quarters[-4:]]
    r4 = [v for v in r4 if v is not None]
    if got != exp or p["pmet_ret"] != (sum(r4) if r4 else None):
        bad.append(k)
check("sector figures equal the Pulse snapshot", not bad, str(bad) if bad else "")
nat = baked["national"]
check("national rate and vacancies equal the Pulse snapshot",
      nat["rate"] == next(r["v"] for r in ds["vacancy_rate_total"]["data"] if r["q"] == q)
      and nat["vac"] == next(r["v"] for r in ds["vacancies_total"]["data"] if r["q"] == q))

# courses vs Landscape latest.json
latest = json.loads((LAND / "data" / "latest.json").read_text())
bad = [k for k in baked["sectors"] if any(
    baked["sectors"][k]["courses"][f] != latest["groups"][k][f] for f in ("n", "providers", "ai", "fee_median"))]
check("course n/providers/ai/fee_median equal Landscape latest.json", not bad, str(bad) if bad else "")
check("national course figures equal latest.json", all(baked["allCourses"][f] == latest["groups"]["all"][f] for f in ("n", "providers", "ai", "fee_median")))
check("top-5 providers equal latest.json", all(baked["sectors"][k]["courses"]["top"] == [t for t in latest["groups"][k]["top"][:5]] for k in baked["sectors"]))
check("course date matches Landscape", baked["asof"]["courses"] == latest["datasets"]["directory"]["updated"][:10])

# top-8 titles vs the raw XLSX
try:
    import openpyxl
    it = openpyxl.load_workbook(LAND / "data" / "raw" / "directory.xlsx", read_only=True).active.iter_rows(values_only=True)
    head = next(it)
    att = {}
    for r in it:
        d = dict(zip(head, r))
        t = re.sub(r"\s+", " ", str(d["coursetitle"] or "").replace("_x000D_", "")).strip()
        try:
            a = float(d["attendancecount"])
        except (TypeError, ValueError):
            continue
        att.setdefault(t, set()).add(int(a))
    bad = [(k, t[0]) for k, c in list(baked["sectors"].items()) + [("all", {"courses": baked["allCourses"]})]
           for t in c["courses"]["titles"] if t[2] not in att.get(t[0], set())]
    check("top-8 course titles exist with those attendance values", not bad, str(bad[:3]) if bad else "")
    check("each sector has up to 8 titles, sorted by attendance", all(
        len(c["courses"]["titles"]) <= 8 and [t[2] for t in c["courses"]["titles"]] == sorted((t[2] for t in c["courses"]["titles"]), reverse=True)
        for c in baked["sectors"].values()))
except FileNotFoundError:
    check("Landscape raw directory.xlsx available for the top-8 check", False)

# SSIC
smap = json.loads((HERE / "data" / "ssic-map.json").read_text())
ssic = json.loads((HERE / "data" / "ssic2025.json").read_text())
keys = set(baked["sectors"]) | {"none"}
check("every SSIC division maps to a sector key or none", set(smap["divisions"]) == set(ssic["divisions"])
      and all(v in keys for v in smap["divisions"].values()))
check("every five-digit code has a division in the map", all(c[:2] in smap["divisions"] for c in ssic["codes"]))
check("holding codes exist in the SSIC 2025 table", all(c in ssic["codes"] for c in smap["holding"]))
check("ambiguous divisions exist in the table", all(d in ssic["divisions"] for d in smap["ambiguous"]))
check("page SSIC table equals ssic2025.json", baked["ssic"] == ssic["codes"])
check("no division mapped to manufacturing outside 10-32", all(d in [f"{i:02d}" for i in range(10, 33)] for d, k in smap["divisions"].items() if k == "manufacturing"))

# shards vs data-findings.md
text = (HERE / "data-findings.md").read_text()
found = dict(re.findall(r"(?:^|\| |: )([A-Z]|Others)(?: \(digit/symbol-first names\))? (d_[0-9a-f]{32})", text.split("Row counts")[0]))
check("27 shard IDs present", len(baked["shards"]) == 27)
check("shard IDs equal data-findings.md", found == baked["shards"], f"{len(found)} parsed")

# page-level rules
low = html.lower()
check("no RP/NACE branding in the page", not re.search(r"republic polytechnic|nace@rp|\bnace\b", re.sub(r'"(sctp|top|titles)":\[\[.*?\]\],?', "", low.split("<script>")[0])) )
hype = ["leverage", "seamless", "unlock", "empower", "robust", "delve"]
body = re.sub(r"const D = .*?;\nconst fmt", "const fmt", html, flags=re.S)
check("no hype words", not [w for w in hype if re.search(r"\b" + w, body.lower())])
check("sessionStorage only used inside try/catch", html.count("sessionStorage") == 2 and "try {" in html)

print(f"\n{len(fails)} failure(s)" if fails else "\nall checks passed")
sys.exit(1 if fails else 0)
